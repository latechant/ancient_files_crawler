# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from pathlib import Path
p = Path(r"D:\AI_Project\ancient_files_crawler\第22组总览表.xlsx")
wb = load_workbook(p)
ws = wb[wb.sheetnames[0]]
for row in ws.iter_rows(min_row=5):
    if row[0].value == "韦广波" and row[1].value == 8429:
        row[5].value = "https://www.mdpi.com/2077-1444/17/3/369"
        row[6].value = "否"
        row[9].value = (
            "2026-09-21 线索：北周建德元年錡馬仁造老君像碑，与北魏正光二年《锜麻仁合家造象记》非同一石。"
            "识典鲁迅辑校石刻手稿仅见后者 https://www.shidianguji.com/book/7563224556832292891/chapter/1lx2cd4f5r19v ；建德元年碑全文录文未取得，故仍记未找到。"
        )
        break
wb.save(p)
print("updated 8429")
