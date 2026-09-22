# -*- coding: utf-8 -*-
"""后台扫描第22组其他人：姜博文 / 马雨衡 / 林靖凯。不碰韦广波进程。"""
from __future__ import annotations

import json
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import requests
from openpyxl import load_workbook

ROOT = Path(r"D:\AI_Project\ancient_files_crawler")
sys.path.insert(0, str(ROOT))

from pipeline.extract_shidian import fetch_html, load_router, book_payload, paragraph_text  # noqa: E402
from pipeline.extract_wikisource import get_wikitext  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")

EXCEL = ROOT / "第22组总览表.xlsx"
LOG = ROOT / "pipeline" / "logs"
CACHE = ROOT / "_cache"
RUN_LOG = LOG / "alt_scan_others.jsonl"
LOG.mkdir(parents=True, exist_ok=True)

OWNERS = ("姜博文", "马雨衡", "林靖凯")
if len(sys.argv) > 1:
    OWNERS = (sys.argv[1],)
    RUN_LOG = LOG / f"alt_scan_{sys.argv[1]}.jsonl"
STAMP = datetime.now().strftime("%Y-%m-%d")
SLEEP = 0.9
MAX_ROUNDS = 6
WIKI_TIMEOUT = 18

HIGH_RE = re.compile(
    r"墓志|墓誌|墓碑|造像|碑|敕|奏|摺|志铭|志銘|寄谕|寄諭|题为|题请|"
    r"府志|县志|縣誌|告示|宗谱|宗譜|公报|公報|金石"
)
STRICT_RE = re.compile(r"寄谕|寄諭|题为|题请|奏报|呈.*批令|相关公文")
SKIP_COPYRIGHT = re.compile(
    r"岳麓书院藏秦简|岭光电|忆往昔|五代十国墓志汇编|西夏李氏世谱|北京档案史料"
)
GENERIC_TITLES = {"执照", "方鼎", "人类", "10月27日", "行书屏", "一百六"}

SESSION = requests.Session()
SESSION.headers.update(
    {
        "User-Agent": "Mozilla/5.0 (compatible; ancient-files-research/1.0; +local archival)",
        "Accept-Language": "zh-CN,zh;q=0.9",
    }
)


def log_event(obj: dict) -> None:
    obj = {"ts": datetime.now().isoformat(timespec="seconds"), **obj}
    with RUN_LOG.open("a", encoding="utf-8") as f:
        f.write(json.dumps(obj, ensure_ascii=False) + "\n")
    print(json.dumps(obj, ensure_ascii=False), flush=True)


def folder_name(seq, title: str) -> str:
    t = str(title or "").strip().strip("《》")
    t = re.sub(r'[\\/:*?"<>|]', "_", t)
    return f"{seq}{t}"


def owner_dirs(owner: str) -> tuple[Path, Path, Path]:
    base = ROOT / owner
    return base / "爬取结果", base / "成果", base / "爬虫总结.txt"


def load_remaining() -> list[dict]:
    wb = load_workbook(EXCEL, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] not in OWNERS or row[2] is None:
            continue
        if str(row[6]) == "是":
            continue
        rows.append(
            {
                "owner": row[0],
                "seq": row[1],
                "title": row[2],
                "url": row[5],
                "publisher": row[7],
                "edition": row[8],
                "note": row[9],
            }
        )
    return rows


def prioritize(rows: list[dict]) -> list[dict]:
    def score(r: dict) -> tuple:
        title = str(r.get("title") or "")
        hi = 0 if HIGH_RE.search(title) else 1
        named = 0 if len(re.findall(r"[\u4e00-\u9fff]", title)) >= 6 else 1
        return (hi, named, str(r["owner"]), int(r["seq"] or 0))

    return sorted(rows, key=score)


def title_tokens(title: str) -> list[str]:
    t = re.sub(r"[《》〈〉（）()〔〕\[\]\s、，。：:；;·\-—]", "", str(title or ""))
    toks = re.findall(r"[\u4e00-\u9fff]{2,8}", t)
    stop = {"相关公文", "情形折", "请旨事", "全文", "原文", "并序"}
    return [x for x in toks if x not in stop][:8]


def queries_for(title: str) -> list[str]:
    raw = str(title or "").strip().strip("《》")
    toks = title_tokens(raw)
    qs = [raw]
    if len(toks) >= 2:
        qs.append("".join(toks[:2]))
    for tok in toks:
        if len(tok) >= 3:
            qs.append(tok)
    out, seen = [], set()
    for q in qs:
        if q and q not in seen:
            seen.add(q)
            out.append(q)
    return out[:3]


def text_matches(title: str, body: str) -> bool:
    t = str(title or "").strip().strip("《》")
    if t in GENERIC_TITLES:
        return False
    if SKIP_COPYRIGHT.search(t + body[:200]):
        return False
    if len(body) < 200:
        return False
    toks = title_tokens(title)
    if not toks or not any(len(tok) >= 3 for tok in toks):
        return False
    hit = sum(1 for tok in toks if tok in body)
    need = 3 if STRICT_RE.search(str(title) or "") else min(2, len(toks))
    return hit >= need and any(len(tok) >= 3 and tok in body for tok in toks)


def walk_hits(obj, acc: list[dict] | None = None) -> list[dict]:
    acc = acc if acc is not None else []
    if isinstance(obj, dict):
        book = obj.get("bookId") or obj.get("book_id")
        chap = obj.get("chapterId") or obj.get("chapter_id")
        name = obj.get("chapterName") or obj.get("bookName") or obj.get("title") or obj.get("name")
        url = obj.get("url") or obj.get("jumpUrl") or ""
        if book and chap:
            acc.append({"url": f"https://www.shidianguji.com/book/{book}/chapter/{chap}", "name": str(name or "")})
        elif isinstance(url, str) and "/book/" in url and "/chapter/" in url:
            acc.append({"url": url, "name": str(name or "")})
        for v in obj.values():
            walk_hits(v, acc)
    elif isinstance(obj, list):
        for it in obj:
            walk_hits(it, acc)
    return acc


def shidian_search(q: str) -> list[dict]:
    url = f"https://www.shidianguji.com/search?query={quote(q)}"
    safe = re.sub(r"[^\w\u4e00-\u9fff]+", "_", q)[:40]
    try:
        html = fetch_html(url, CACHE / "shidian_search" / f"{safe}.html")
        router = load_router(html)
    except Exception as exc:
        log_event({"event": "search_fail", "q": q[:40], "error": type(exc).__name__})
        return []
    uniq, seen = [], set()
    for h in walk_hits(router):
        if h["url"] in seen:
            continue
        seen.add(h["url"])
        uniq.append(h)
    return uniq[:10]


def wiki_search(q: str) -> list[dict]:
    try:
        r = SESSION.get(
            "https://zh.wikisource.org/w/api.php",
            params={"action": "query", "list": "search", "srsearch": q, "srlimit": 6, "format": "json"},
            timeout=WIKI_TIMEOUT,
        )
        r.raise_for_status()
        return [{"title": it.get("title") or ""} for it in r.json().get("query", {}).get("search", [])]
    except Exception as exc:
        log_event({"event": "wiki_fail", "q": q[:40], "error": type(exc).__name__})
        return []


def fetch_shidian_chapter(url: str) -> str:
    m = re.search(r"/book/([^/]+)/chapter/([^/?#]+)", url)
    if not m:
        return ""
    book_id, cid = m.group(1), m.group(2)
    html = fetch_html(url, CACHE / "shidian" / f"{book_id}_{cid}.html")
    payload = book_payload(load_router(html))
    paras = payload.get("paragraphList") or []
    return "\n".join(t for t in (paragraph_text(p).strip() for p in paras) if t)


def save_wb(wb, path: Path, tries: int = 12) -> None:
    last = None
    for i in range(tries):
        try:
            wb.save(path)
            return
        except (OSError, PermissionError, EOFError) as exc:
            last = exc
            time.sleep(1.4 + i * 0.5)
    raise last  # type: ignore[misc]


def dump_found(owner: str, seq, title: str, url: str, body: str, note: str) -> None:
    result, product, _ = owner_dirs(owner)
    header = f"# {title}\n# 来源: {url}\n# 编码: UTF-8\n# 归档日期: {STAMP}\n\n"
    content = header + body.strip() + "\n"
    folder = result / folder_name(seq, title)
    prod = product / str(title).strip("《》")
    folder.mkdir(parents=True, exist_ok=True)
    prod.mkdir(parents=True, exist_ok=True)
    (folder / "正文.txt").write_text(content, encoding="utf-8")
    (prod / "正文.txt").write_text(content, encoding="utf-8")
    (folder / "来源网址.txt").write_text(url + "\n", encoding="utf-8")
    (prod / "来源网址.txt").write_text(url + "\n", encoding="utf-8")
    meta = {
        "owner": owner,
        "seq": seq,
        "title": title,
        "url": url,
        "note": note,
        "saved_at": STAMP,
        "ok": True,
        "chars": len(body),
    }
    blob = json.dumps(meta, ensure_ascii=False, indent=2)
    (folder / "基本信息.json").write_text(blob, encoding="utf-8")
    (prod / "基本信息.json").write_text(blob, encoding="utf-8")


def mark_excel(found: list[dict]) -> int:
    if not found:
        return 0
    wb = load_workbook(EXCEL)
    ws = wb[wb.sheetnames[0]]
    by = {(it["owner"], int(it["seq"])): it for it in found}
    n = 0
    for row in ws.iter_rows(min_row=5):
        key = (row[0].value, row[1].value)
        if key not in by:
            continue
        it = by[key]
        row[5].value = it["url"]
        row[6].value = "是"
        row[7].value = it.get("publisher") or "不详"
        row[8].value = it.get("edition") or "公开录文页"
        old = str(row[9].value or "").strip()
        note = it.get("note") or f"{STAMP} 后台扫描核验归档。"
        row[9].value = note if old in ("", "未找到", "无") else f"{old}；{note}"
        n += 1
    save_wb(wb, EXCEL)
    return n


def count_status() -> dict[str, tuple[int, int]]:
    wb = load_workbook(EXCEL, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out = {o: [0, 0] for o in OWNERS}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] not in OWNERS or row[2] is None:
            continue
        if str(row[6]) == "是":
            out[row[0]][0] += 1
        else:
            out[row[0]][1] += 1
    return {k: (v[0], v[1]) for k, v in out.items()}


def update_summary(found: list[dict], reason: str) -> None:
    by_owner: dict[str, int] = {}
    for it in found:
        by_owner[it["owner"]] = by_owner.get(it["owner"], 0) + 1
    status = count_status()
    for owner, extra in by_owner.items():
        _, summary = owner_dirs(owner)[0], owner_dirs(owner)[2]
        yes, no = status.get(owner, (0, 0))
        line = f"{STAMP} 组内异源扫描：{reason}。本进程新归档 {extra} 条。当前找到{yes}条，未找到{no}条。\n"
        old = summary.read_text(encoding="utf-8") if summary.exists() else ""
        old = re.sub(r"找到\d+条，未找到\d+条。", f"找到{yes}条，未找到{no}条。", old, count=1)
        summary.write_text(line + old, encoding="utf-8")
    log_event({"event": "summary", "reason": reason, "by_owner": by_owner, "status": {k: list(v) for k, v in status.items()}})


def try_archive(rec: dict, url: str, body: str, seen: set[tuple]) -> dict | None:
    key = (rec["owner"], int(rec["seq"]))
    if key in seen:
        return None
    if not text_matches(str(rec["title"]), body):
        log_event({"event": "mismatch", "owner": rec["owner"], "seq": rec["seq"], "url": url, "chars": len(body)})
        return None
    note = f"{STAMP} 组内异源扫描核验，已保存匹配录文。"
    dump_found(rec["owner"], rec["seq"], str(rec["title"]).strip("《》"), url, body, note)
    seen.add(key)
    hit = {
        "owner": rec["owner"],
        "seq": int(rec["seq"]),
        "title": str(rec["title"]).strip("《》"),
        "url": url,
        "note": note,
    }
    log_event({"event": "archived", **hit, "chars": len(body)})
    return hit


def scan_cache(rows: list[dict], seen: set[tuple]) -> list[dict]:
    found = []
    cache_dir = CACHE / "shidian"
    if not cache_dir.exists():
        return found
    needles = []
    for rec in rows:
        for tok in title_tokens(str(rec["title"])):
            if len(tok) >= 3:
                needles.append((tok, rec))
    files = list(cache_dir.glob("*.html"))
    log_event({"event": "cache_scan", "files": len(files), "needles": len(needles)})
    for fp in files:
        try:
            html = fp.read_text(encoding="utf-8")
            if "window._ROUTER_DATA" not in html:
                continue
            payload = book_payload(load_router(html))
            paras = payload.get("paragraphList") or []
            body = "\n".join(t for t in (paragraph_text(p).strip() for p in paras) if t)
        except Exception:
            continue
        if len(body) < 200:
            continue
        m = re.search(r"([A-Za-z0-9]+)_([A-Za-z0-9]+)\.html$", fp.name)
        url = f"https://www.shidianguji.com/book/{m.group(1)}/chapter/{m.group(2)}" if m else str(fp)
        for tok, rec in needles:
            if (rec["owner"], int(rec["seq"])) in seen:
                continue
            if tok in body:
                hit = try_archive(rec, url, body, seen)
                if hit:
                    found.append(hit)
    return found


def scan_item(rec: dict, seen: set[tuple]) -> list[dict]:
    title = str(rec["title"])
    if title.strip().strip("《》") in GENERIC_TITLES or SKIP_COPYRIGHT.search(title):
        return []
    found = []
    url0 = str(rec.get("url") or "")
    cands: list[tuple[str, str]] = []
    if "shidianguji.com" in url0 and "/chapter/" in url0:
        cands.append((url0, "shidian"))
    for q in queries_for(title):
        for h in shidian_search(q):
            cands.append((h["url"], "shidian"))
        time.sleep(SLEEP)
        for h in wiki_search(q):
            cands.append((h["title"], "wiki"))
        time.sleep(0.3)
    seen_u = set()
    for target, kind in cands:
        if target in seen_u:
            continue
        seen_u.add(target)
        try:
            if kind == "shidian":
                body = fetch_shidian_chapter(target)
                url = target
            else:
                body = get_wikitext(target) or ""
                url = "https://zh.wikisource.org/wiki/" + quote(target)
        except Exception as exc:
            log_event({"event": "fetch_fail", "owner": rec["owner"], "seq": rec["seq"], "error": type(exc).__name__})
            continue
        if not body:
            continue
        hit = try_archive(rec, url, body, seen)
        if hit:
            found.append(hit)
            break
        time.sleep(0.2)
    return found


def main() -> None:
    log_event({"event": "start", "owners": list(OWNERS), "max_rounds": MAX_ROUNDS})
    seen: set[tuple] = set()
    total: list[dict] = []
    rows = prioritize(load_remaining())
    log_event({"event": "loaded", "n": len(rows), "status": {k: list(v) for k, v in count_status().items()}})
    cache_hits = scan_cache(rows, seen)
    if cache_hits:
        mark_excel(cache_hits)
        update_summary(cache_hits, "缓存回扫")
        total.extend(cache_hits)
    for rnd in range(1, MAX_ROUNDS + 1):
        rows = prioritize(load_remaining())
        log_event({"event": "round", "n": rnd, "remaining": len(rows)})
        batch = []
        for rec in rows:
            key = (rec["owner"], int(rec["seq"]))
            if key in seen:
                continue
            try:
                batch.extend(scan_item(rec, seen))
            except Exception as exc:
                log_event({"event": "item_fail", "owner": rec["owner"], "seq": rec["seq"], "error": type(exc).__name__})
            time.sleep(0.15)
        if batch:
            mark_excel(batch)
            update_summary(batch, f"第{rnd}轮")
            total.extend(batch)
        log_event({"event": "round_done", "n": rnd, "new": len(batch), "seqs": [(x["owner"], x["seq"]) for x in batch]})
        time.sleep(6)
    log_event({"event": "end", "new": len(total)})


if __name__ == "__main__":
    main()
