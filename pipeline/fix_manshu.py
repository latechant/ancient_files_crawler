# -*- coding: utf-8 -*-
"""补抽维基文库《蠻書（四庫全書本）》并回写总览表第3815条。"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"D:\AI_Project\ancient_files_crawler")
sys.path.insert(0, str(ROOT))
from pipeline.extract_wikisource import extract_prefixed
from pipeline.run_wei_batch import archive_texts, write_meta, write_sources, EXCEL, RESULT, PRODUCT

sys.stdout.reconfigure(encoding="utf-8")
STAMP = datetime.now().strftime("%Y-%m-%d")

data = extract_prefixed("蠻書 (四庫全書本)")
print("pages", data["titles"])
texts = {}
for title, body in data["texts"].items():
    if not body.strip():
        continue
    short = title.replace("蠻書 (四庫全書本)/", "蠻書_").replace("蠻書 (四庫全書本)", "蠻書_目錄")
    texts[short] = body
    print("got", short, "chars", len(body))

rec = {
    "seq": 3815,
    "title": "云南记",
    "folder": "3815云南记",
    "publisher": "武英殿聚珍版/四库全书本",
    "edition": "唐樊绰《蛮书》十卷，亦名云南志、云南史记；袁滋《云南记》已佚。本条归档四库本《蛮书》全文作别名核验。",
    "note": (
        f"{STAMP} 袁滋《云南记》原书已佚。"
        "现存可核全文为樊绰《蠻書（四庫全書本）》十卷（别名云南志/云南史记，部分文献亦称云南记）。"
        "已按卷次自维基文库提取，不能径视为袁滋原书。"
    ),
}
folder = RESULT / rec["folder"]
product = PRODUCT / rec["title"]
n = archive_texts(folder, product, texts)
sources = [
    "https://zh.wikisource.org/wiki/蠻書_(四庫全書本)",
    "https://zh.wikisource.org/wiki/蠻書_(四庫全書本)/全覽",
    "https://zh.wikipedia.org/wiki/云南志",
]
write_sources(folder / "来源网址.txt", sources)
write_sources(product / "来源网址.txt", sources)
(folder / "判定说明.txt").write_text(rec["note"] + "\n", encoding="utf-8")
(product / "判定说明.txt").write_text(rec["note"] + "\n", encoding="utf-8")
meta = {"source": sources[0], "saved_txt": n, "ok": n > 0, "url": sources[0], "pages": data["titles"]}
write_meta(folder / "基本信息.json", rec, meta)
write_meta(product / "基本信息.json", rec, meta)
print("saved_txt", n)

wb = load_workbook(EXCEL)
ws = wb[wb.sheetnames[0]]
for row in ws.iter_rows(min_row=5):
    if row[0].value == "韦广波" and row[1].value == 3815:
        row[5].value = sources[0]
        row[6].value = "是" if n else "否"
        row[7].value = rec["publisher"]
        row[8].value = rec["edition"]
        row[9].value = rec["note"]
        break
wb.save(EXCEL)
print("excel ok")
