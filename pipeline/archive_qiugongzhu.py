# -*- coding: utf-8 -*-
"""归档宋仇公著墓志：母项《八琼室金石补正续编》卷三十二录文。"""
from __future__ import annotations

import json
import sys
import time
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(r"D:\AI_Project\ancient_files_crawler")
RESULT = ROOT / "韦广波" / "爬取结果"
PRODUCT = ROOT / "韦广波" / "成果"
EXCEL = ROOT / "第22组总览表.xlsx"
STAMP = datetime.now().strftime("%Y-%m-%d")
SUMMARY = ROOT / "韦广波" / "爬虫总结.txt"

SEQ = 14071
TITLE = "宋仇公著墓志"
FOLDER = "14071宋仇公著墓志"
FILE = "八琼室金石补正续编_卷三十二_仇公著墓志铭.txt"
URL = "https://www.shidianguji.com/book/7549461577028599854/chapter/1lcvz8ltnwv0q"

raw = (ROOT / "pipeline" / "logs" / "_qiugongzhu.txt").read_text(encoding="utf-8")
# 只取墓志铭本体，不含后续龙兴寺等条。
lines = raw.splitlines()
kept = []
for line in lines:
    if kept and ("龙兴寺大悲阁记" in line or "龍興寺大悲閣記" in line):
        # 同行若拼接了下一条，截断。
        i = line.find("龙兴寺")
        if i < 0:
            i = line.find("龍興寺")
        if i > 0:
            kept.append(line[:i].rstrip())
        break
    kept.append(line)
BODY = "\n".join(kept).strip()
NOTE = (
    f"{STAMP} 条目为宋仇公著墓志。母项《八琼室金石补正续编》卷三十二"
    "《仇公著墓志铭》已从识典抽出全文。柳子文撰、王同老书、晁端德篆；"
    "绍圣三年四月卒，十月葬青州云门里。跋另见《金石三编》。"
)


def dump() -> None:
    header = f"# {TITLE}\n# 来源: {URL}\n# 编码: UTF-8\n# 归档日期: {STAMP}\n\n"
    content = header + BODY + "\n"
    folder = RESULT / FOLDER
    product = PRODUCT / TITLE
    folder.mkdir(parents=True, exist_ok=True)
    product.mkdir(parents=True, exist_ok=True)
    (folder / FILE).write_text(content, encoding="utf-8")
    (product / FILE).write_text(content, encoding="utf-8")
    (folder / "来源网址.txt").write_text(URL + "\n", encoding="utf-8")
    (product / "来源网址.txt").write_text(URL + "\n", encoding="utf-8")
    meta = {
        "seq": SEQ,
        "title": TITLE,
        "url": URL,
        "publisher": "不详",
        "edition": "清陆增祥《八琼室金石补正续编》卷三十二；识典本。",
        "note": NOTE,
        "saved_at": STAMP,
        "ok": True,
        "chars": len(BODY),
    }
    (folder / "基本信息.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    (product / "基本信息.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    print("saved files", SEQ, len(BODY))


def mark_excel() -> None:
    last = None
    for i in range(10):
        try:
            wb = load_workbook(EXCEL)
            ws = wb[wb.sheetnames[0]]
            for row in ws.iter_rows(min_row=5):
                if row[0].value != "韦广波":
                    continue
                if row[1].value != SEQ:
                    continue
                row[5].value = URL
                row[6].value = "是"
                row[7].value = "不详"
                row[8].value = "清陆增祥《八琼室金石补正续编》卷三十二；识典本。"
                old = str(row[9].value or "").strip()
                row[9].value = NOTE if old in ("", "未找到", "无") else f"{old}；{NOTE}"
            wb.save(EXCEL)
            print("excel ok")
            return
        except Exception as exc:
            last = exc
            print("excel retry", i, type(exc).__name__)
            time.sleep(1.5)
    raise last  # type: ignore[misc]


def update_summary() -> None:
    extra = (
        f"\n{STAMP} 改查识典/维基/实录母项：新增找到 14071宋仇公著墓志"
        f"（八琼室金石补正续编卷三十二，{len(BODY)}字）。\n"
    )
    if SUMMARY.exists():
        text = SUMMARY.read_text(encoding="utf-8")
        if "14071宋仇公著墓志" not in text:
            SUMMARY.write_text(text.rstrip() + extra, encoding="utf-8")
    print("summary ok")


if __name__ == "__main__":
    if len(BODY) < 400:
        raise SystemExit(f"body too short: {len(BODY)}")
    dump()
    mark_excel()
    update_summary()
