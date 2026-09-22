# -*- coding: utf-8 -*-
import json
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
p = r"D:\AI_Project\ancient_files_crawler\第22组总览表.xlsx"
wb = load_workbook(p, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = []
for row in ws.iter_rows(min_row=5, values_only=True):
    if row[0] != "韦广波" or row[2] is None:
        continue
    if str(row[6]) == "是":
        continue
    rows.append({
        "seq": row[1],
        "title": row[2],
        "url": row[5],
        "publisher": row[7],
        "edition": row[8],
        "note": (str(row[9] or ""))[:180],
    })
print("remaining", len(rows))
for r in rows[:40]:
    print(r["seq"], r["title"], r["url"])
out = r"D:\AI_Project\ancient_files_crawler\韦广波\_remaining.json"
with open(out, "w", encoding="utf-8") as f:
    json.dump(rows, f, ensure_ascii=False, indent=2)
print("wrote", out)
