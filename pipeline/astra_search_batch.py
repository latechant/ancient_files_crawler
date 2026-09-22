# -*- coding: utf-8 -*-
"""用 aizex gpt-6-astra 为未找到条目生成检索方案（关键词/别名/母项/候选网址）。"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

APM = Path(r"D:\AI_Project\automatic_problem_maker")
ROOT = Path(r"D:\AI_Project\ancient_files_crawler")
sys.path.insert(0, str(APM))

from tools.llm_client import call_llm  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")

EXCEL = ROOT / "第22组总览表.xlsx"
LOG = ROOT / "pipeline" / "logs"
LOG.mkdir(parents=True, exist_ok=True)

SYSTEM = """你是古籍与近代文献电子资源检索顾问。任务是给清单条目生成可核验的检索方案，不是直接宣称已找到全文。
规则：
1. 不要只搜完整题名；给出简繁体、别名、作者、版本、母项（文集/方志/报刊/档案汇编）与子项定位线索。
2. 优先公开可核来源：识典古籍、维基文库、中国哲学书电子化计划、国家图书馆、CADAL、Commons。
3. 现代期刊、2015 年后整理本、仍在版权期的著作：只给书目/馆藏/DOI 线索，标明不可整本下载。
4. 找不到就明确写「未找到」，不要编造能打开的全文 URL。
5. 只输出 JSON 数组，不要 markdown。每项字段：
seq, title, item_type, keywords(数组), aliases(数组), mother_works(数组),
candidate_urls(数组，每项含 url, why, likely_fulltext true/false),
verdict(找到候选/仅线索/未找到), notes
"""


def remaining(limit: int = 10, offset: int = 0) -> list[dict]:
    wb = load_workbook(EXCEL, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    skipped = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] != "韦广波" or row[6] == "是" or row[2] is None:
            continue
        if skipped < offset:
            skipped += 1
            continue
        rows.append(
            {
                "seq": row[1],
                "title": row[2],
                "url": row[5],
                "publisher": row[7],
                "edition": row[8],
                "note": row[9],
            }
        )
        if len(rows) >= limit:
            break
    return rows


def main() -> None:
    items = remaining(12, offset=10)
    print("batch", len(items), [i["seq"] for i in items])
    user = (
        "请为下列韦广波未找到条目生成检索方案。禁止只给国图/CADAL/维基文库首页。"
        "必须给出尽量具体的书名卷次、期刊年期或识典古籍/维基文库章节 URL。"
        "找不到具体页就 verdict=未找到。清单很多不是独立书籍。\n"
        + json.dumps(items, ensure_ascii=False, indent=2)
    )
    text = call_llm(
        user,
        system_prompt=SYSTEM,
        temperature=0.2,
        timeout_seconds=180,
        model_name_override="gpt-6-astra",
        max_tokens=4000,
    )
    raw_path = LOG / f"astra_search_{datetime.now().strftime('%Y%m%d_%H%M%S')}.raw.txt"
    raw_path.write_text(text or "", encoding="utf-8")
    print("raw", raw_path, "chars", len(text or ""))
    blob = (text or "").strip()
    if blob.startswith("```"):
        blob = blob.strip("`")
        if blob.startswith("json"):
            blob = blob[4:]
        blob = blob.strip()
    start = blob.find("[")
    end = blob.rfind("]")
    parsed = json.loads(blob[start : end + 1]) if start >= 0 and end > start else []
    out = LOG / "astra_search_latest.json"
    out.write_text(json.dumps(parsed, ensure_ascii=False, indent=2), encoding="utf-8")
    print("parsed_items", len(parsed))
    for rec in parsed:
        print(
            rec.get("seq"),
            rec.get("verdict"),
            rec.get("title"),
            len(rec.get("candidate_urls") or []),
            "urls",
        )


if __name__ == "__main__":
    main()
