# -*- coding: utf-8 -*-
"""归档苏日荣墓志录文（书法资料网转引自《唐代墓誌彙編》）。"""
from __future__ import annotations

import json
import sys
import time
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(r"D:\AI_Project\ancient_files_crawler")
RESULT = ROOT / "韦广波" / "爬取结果"
PRODUCT = ROOT / "韦广波" / "成果"
EXCEL = ROOT / "第22组总览表.xlsx"
STAMP = datetime.now().strftime("%Y-%m-%d")

SEQ = 16632
TITLE = "苏日荣墓志"
FOLDER = "16632苏日荣墓志"
FILE = "唐故特进行虔王傅扶风县开国伯上柱国兼英武军右厢兵马使苏公墓志铭.txt"
URL = "https://shufaziliao.com/maobi/6376a0b8ac2257e6190a777a759602e8.html"
BODY = """唐故特進行虔王傅扶風縣開國伯上柱國兼英武軍右廂兵馬使蘇公墓誌銘并序
將仕郎守秘書省校書郎房次卿撰

貞元十四年六月廿九日，文武之上將，社稷之忠臣，特進行虔王傅兼英武軍右廂兵馬使蘇公薨于位。皇帝輟食去樂，賵賻加等。公諱日榮，字德昌，京兆武功人也。得姓受氏，著經見史；孝家忠國，角立傑出。曾祖德瑶，中散大夫、翼州刺史；祖知廉，銀青光禄大夫、甘州刺史贈尚書左僕射；並積善種德，餘慶垂休，鍾秀於公，公克負荷。天寶初，以一子出身授安鄉郡大夏縣丞，文理優柔，縣境和洽，時無軍旅之事，方脩俎豆之儀。無何，禄山始亂，四海兵興，於是慨然有投筆之心，枕戈之志。時客遊靜邊軍，乃陰結豪猛，潛期殄除。至一暮夜，義徒竊發，遂斬逆將安守一、周萬頃以甘心焉。玄宗異之，特遷右清道率，充振武軍副使。至德初，領朔方三郡之士，衛北極九重之嚴，奮千人被練之師，殲九蕃同羅之衆，拜右領軍衛將軍。宇宙載寧，乘輿返正，轉本衛大將軍。蠻夷亂華，天子巡陝，率紀綱之僕，為腹心之臣，封扶風縣開國子，除右千牛衛大將軍。主上龍飛，録功班爵，當監撫之日，有調護之勳，改右武衛大將軍，充仗內教坊使。六師無闕，八音克諧，寵榮一時，熏灼四海。過有小譴，情實寡尤，貶萬州司馬。跼蹐惟省，精誠感通，俄而優詔追還，恩遇如舊，遷右監門衛將軍，遽昇本衛大將軍，又歷韶王傅，復用虔王傅兼英武軍右廂兵馬使。方嘉節制之能，將付旄鉞之任，而天玄莫問，神寂奚徵，知命順時，即世從化。時年七十有九。嗣子審、宙、定等，以貞元十四年八月七日泣血徒跣，扶護靈輀，歸葬于東都河南縣金谷原，安定郡夫人智氏祔焉，禮也。公幼而好學，事親竭力；長而尚氣，許國亡軀。身已貴高，接下之禮有餘；節彌清亮，□人之急不倦。居常秉金玉之度，臨敵奮貔虎之威。德著有不勝之容，功多無自矜之色。故得久承密恩，雖疾轉任，獨蒙殊渥，既殁猶思加以鑒裁知人，推引得士。辛太原之豪傑，事公以師儀；渾河中之功名，敬公以兄禮。其餘茅土之俊，公卿之妙，無不昇堂入室，結友定交。嗚呼！人之云亡，邦國殄瘁。成蹊之感，復睹斯時；大樹之悲，寧聞往日。著金石之悠久，備陵谷之遷移，乃為銘曰：

功名開國復承家，崇堂邃宇延綺羅，賢豪英達來經過，前妙舞兮後清歌，愁時少兮歡樂多，一朝棄捐託山阿，聖智勇力如之何！惟玆刊石流惠和。嗣子進士□□
"""
NOTE = (
    f"{STAMP} 核为《唐故特进行虔王傅扶风县开国伯上柱国兼英武军右厢兵马使苏公墓志铭》，"
    "房次卿撰，贞元十四年。母项《唐代墓誌彙編》下；千唐志斋藏石。原表书法网仅为拓片页，本轮改用录文页归档。"
)


def dump() -> None:
    header = f"# {TITLE}\n# 来源: {URL}\n# 编码: UTF-8\n# 归档日期: {STAMP}\n\n"
    content = header + BODY.strip() + "\n"
    folder = RESULT / FOLDER
    product = PRODUCT / TITLE
    folder.mkdir(parents=True, exist_ok=True)
    product.mkdir(parents=True, exist_ok=True)
    (folder / FILE).write_text(content, encoding="utf-8")
    (product / FILE).write_text(content, encoding="utf-8")
    (folder / "来源网址.txt").write_text(URL + "\n", encoding="utf-8")
    (product / "来源网址.txt").write_text(URL + "\n", encoding="utf-8")
    meta = {
        "seq": SEQ,
        "title": TITLE,
        "url": URL,
        "publisher": "上海古籍出版社",
        "edition": "周绍良《唐代墓誌彙編》下；千唐志斋藏石。唐贞元十四年，墓主苏日荣字德昌。",
        "note": NOTE,
        "saved_at": STAMP,
        "ok": True,
    }
    (folder / "基本信息.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    (product / "基本信息.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    print("saved files", SEQ)


def mark_excel() -> None:
    last = None
    for i in range(10):
        try:
            wb = load_workbook(EXCEL)
            ws = wb[wb.sheetnames[0]]
            for row in ws.iter_rows(min_row=5):
                if row[0].value != "韦广波":
                    continue
                if row[1].value != SEQ:
                    continue
                row[5].value = URL
                row[6].value = "是"
                row[7].value = "上海古籍出版社"
                row[8].value = "周绍良《唐代墓誌彙編》下；千唐志斋藏石。唐贞元十四年。"
                old = str(row[9].value or "").strip()
                row[9].value = NOTE if old in ("", "未找到", "无") else f"{old}；{NOTE}"
            wb.save(EXCEL)
            print("excel ok")
            return
        except Exception as exc:
            last = exc
            print("excel retry", i, type(exc).__name__)
            time.sleep(1.5)
    raise last  # type: ignore[misc]


if __name__ == "__main__":
    dump()
    mark_excel()
