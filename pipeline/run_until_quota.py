# -*- coding: utf-8 -*-
"""循环调用 aizex gpt-6-astra，直到额度耗尽或没有可再挖的条目。"""
from __future__ import annotations

import json
import re
import sys
import time
from datetime import datetime
from html import unescape
from pathlib import Path
from urllib.parse import unquote, urlparse

import requests
from openpyxl import load_workbook

APM = Path(r"D:\AI_Project\automatic_problem_maker")
ROOT = Path(r"D:\AI_Project\ancient_files_crawler")
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(APM))

from src.llm.llm_client import LlmQuotaExhaustedError, call_llm, is_llm_quota_exhausted_error  # noqa: E402

from pipeline.extract_shidian import (  # noqa: E402
    book_payload,
    fetch_html,
    load_router,
    paragraph_text,
)
from pipeline.extract_wikisource import get_wikitext  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")

OWNER = sys.argv[1] if len(sys.argv) > 1 else "韦广波"
EXCEL = ROOT / "第22组总览表.xlsx"
RESULT = ROOT / OWNER / "爬取结果"
PRODUCT = ROOT / OWNER / "成果"
LOG = ROOT / "pipeline" / "logs"
CACHE = ROOT / "_cache"
LOG.mkdir(parents=True, exist_ok=True)
STAMP = datetime.now().strftime("%Y-%m-%d")
RUN_LOG = LOG / f"until_quota_{OWNER}.jsonl"
SUMMARY = ROOT / OWNER / "爬虫总结.txt"

BATCH_SIZE = 8
TIMEOUT = 180
MAX_TOKENS = 4000
SLEEP_BETWEEN = 2.0

SYSTEM = """你是古籍与近代文献电子资源检索顾问。任务是给清单条目生成可核验的检索方案，不是直接宣称已找到全文。
规则：
1. 不要只搜完整题名；给出简繁体、别名、作者、版本、母项（文集/方志/报刊/档案汇编）与子项定位线索。
2. 优先公开可核来源：识典古籍、维基文库、书法网/墓志录文站、Commons。
3. 现代期刊、2015 年后整理本、仍在版权期的著作：只给书目/馆藏/DOI 线索，标明不可整本下载。
4. 找不到就明确写「未找到」，不要编造能打开的全文 URL。禁止国图/CADAL/维基文库/识典首页。
5. 只输出 JSON 数组，不要 markdown。每项字段：
seq, title, item_type, keywords(数组), aliases(数组), mother_works(数组),
candidate_urls(数组，每项含 url, why, likely_fulltext true/false),
verdict(找到候选/仅线索/未找到), notes
"""

USER_HEAD = (
    f"请为下列{OWNER}未找到条目生成检索方案。禁止只给国图/CADAL/维基文库首页。"
    "必须给出尽量具体的书名卷次、期刊年期或识典古籍/维基文库章节 URL。"
    "找不到具体页就 verdict=未找到。清单很多不是独立书籍。\n"
)

DEEP_HEAD = (
    "这些条目此前只有笼统线索、仍未取得公开全文。请再检索一轮："
    "只返回更具体的识典章节、维基文库卷次、墓志录文页或方志卷 URL。"
    "不要重复首页。找不到就 verdict=未找到。\n"
)

HIGH_RE = re.compile(
    r"墓志|墓誌|墓碑|造像|碑|敕|奏|摺|志铭|志銘|简报|清理|发掘|發掘|"
    r"府志|县志|縣誌|厅志|廳志|汇编|彙編|册府|金石|宗谱|宗譜|告示"
)

FETCH_HOSTS = (
    "shidianguji.com",
    "zh.wikisource.org",
    "wikisource.org",
    "shufawang.cn",
    "www.shufawang.cn",
    "xiaoxiong360.com",
    "ybsfzx.com",
    "www.ybsfzx.com",
    "livingbuddha.cn",
    "www.livingbuddha.cn",
)

SESSION = requests.Session()
SESSION.headers.update(
    {
        "User-Agent": "Mozilla/5.0 (compatible; ancient-files-research/1.0; +local archival)",
        "Accept-Language": "zh-CN,zh;q=0.9",
    }
)


def log_event(obj: dict) -> None:
    obj = {"ts": datetime.now().isoformat(timespec="seconds"), **obj}
    with RUN_LOG.open("a", encoding="utf-8") as f:
        f.write(json.dumps(obj, ensure_ascii=False) + "\n")
    print(json.dumps(obj, ensure_ascii=False), flush=True)


def save_wb(wb, path: Path, tries: int = 8) -> None:
    last = None
    for i in range(tries):
        try:
            wb.save(path)
            return
        except (OSError, PermissionError, EOFError) as exc:
            last = exc
            time.sleep(1.2 + i * 0.4)
    raise last  # type: ignore[misc]


def folder_name(seq, title: str) -> str:
    t = str(title or "").strip().strip("《》")
    t = re.sub(r'[\\/:*?"<>|]', "_", t)
    return f"{seq}{t}"


def has_plan(seq, title: str, note: str) -> bool:
    if "gpt-6-astra" in (note or ""):
        return True
    folder = RESULT / folder_name(seq, title)
    return (folder / "本轮补充" / "astra检索方案.json").exists()


def load_remaining(only_unplanned: bool = True) -> list[dict]:
    wb = load_workbook(EXCEL, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] != OWNER or row[2] is None:
            continue
        if str(row[6]) == "是":
            continue
        rec = {
            "seq": row[1],
            "title": row[2],
            "url": row[5],
            "publisher": row[7],
            "edition": row[8],
            "note": row[9],
        }
        if only_unplanned and has_plan(rec["seq"], rec["title"], str(rec["note"] or "")):
            continue
        rows.append(rec)
    return rows


def prioritize(rows: list[dict]) -> list[dict]:
    def score(r: dict) -> tuple:
        title = str(r.get("title") or "")
        hi = 0 if HIGH_RE.search(title) else 1
        named = 0 if len(re.findall(r"[\u4e00-\u9fff]", title)) >= 6 else 1
        return (hi, named, int(r["seq"] or 0))

    return sorted(rows, key=score)


def parse_json_array(text: str) -> list:
    blob = (text or "").strip()
    if blob.startswith("```"):
        blob = blob.strip("`")
        if blob.startswith("json"):
            blob = blob[4:]
        blob = blob.strip()
    start = blob.find("[")
    end = blob.rfind("]")
    if start < 0 or end <= start:
        return []
    return json.loads(blob[start : end + 1])


def call_astra(items: list[dict], head: str) -> list:
    user = head + json.dumps(items, ensure_ascii=False, indent=2)
    text = call_llm(
        user,
        system_prompt=SYSTEM,
        temperature=0.2,
        timeout_seconds=TIMEOUT,
        model_name_override="gpt-6-astra",
        max_tokens=MAX_TOKENS,
        fallback_models=(),
        use_cache=False,
    )
    raw_path = LOG / f"astra_search_{datetime.now().strftime('%Y%m%d_%H%M%S')}.raw.txt"
    raw_path.write_text(text or "", encoding="utf-8")
    parsed = parse_json_array(text or "")
    latest = LOG / "astra_search_latest.json"
    latest.write_text(json.dumps(parsed, ensure_ascii=False, indent=2), encoding="utf-8")
    log_event(
        {
            "event": "astra_batch",
            "n_items": len(items),
            "seqs": [i["seq"] for i in items],
            "raw": str(raw_path.name),
            "parsed": len(parsed),
            "chars": len(text or ""),
        }
    )
    for rec in parsed:
        print(
            rec.get("seq"),
            rec.get("verdict"),
            rec.get("title"),
            len(rec.get("candidate_urls") or []),
            "urls",
            flush=True,
        )
    return parsed


def apply_plans(plans: list) -> int:
    if not plans:
        return 0
    wb = load_workbook(EXCEL)
    ws = wb[wb.sheetnames[0]]
    by_seq = {}
    for p in plans:
        try:
            by_seq[int(p["seq"])] = p
        except (TypeError, ValueError, KeyError):
            continue
    n = 0
    for row in ws.iter_rows(min_row=5):
        if row[0].value != OWNER:
            continue
        seq = row[1].value
        if seq not in by_seq:
            continue
        p = by_seq[seq]
        kws = "、".join((p.get("keywords") or [])[:6])
        urls = "; ".join(u.get("url") for u in (p.get("candidate_urls") or [])[:3] if u.get("url"))
        note = (
            f"{STAMP} gpt-6-astra：{p.get('verdict')}。"
            f"类型={p.get('item_type')}。关键词={kws}。"
            f"母项={('、'.join(p.get('mother_works') or [])[:80])}。"
            f"候选站={urls}。{p.get('notes') or ''}"
        )
        old = str(row[9].value or "").strip()
        new_full = [
            u.get("url")
            for u in (p.get("candidate_urls") or [])
            if u.get("url") and u.get("likely_fulltext")
        ]
        if "gpt-6-astra" in old and not new_full and p.get("verdict") == "未找到":
            pass
        else:
            row[9].value = note if old in ("", "未找到", "无") else f"{old}；{note}"
        n += 1
        folder = RESULT / folder_name(seq, p.get("title") or row[2].value)
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "本轮补充").mkdir(exist_ok=True)
        (folder / "本轮补充" / "astra检索方案.json").write_text(
            json.dumps(p, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        url_lines = [u.get("url") for u in (p.get("candidate_urls") or []) if u.get("url")]
        if url_lines:
            (folder / "来源网址.txt").write_text("\n".join(url_lines) + "\n", encoding="utf-8")
    save_wb(wb, EXCEL)
    log_event({"event": "apply_notes", "n": n})
    return n


def host_ok(url: str) -> bool:
    try:
        host = urlparse(url).netloc.lower()
    except Exception:
        return False
    return any(host == h or host.endswith("." + h) for h in FETCH_HOSTS)


def html_to_text(html: str) -> str:
    html = re.sub(r"(?is)<script[\s\S]*?</script>", " ", html)
    html = re.sub(r"(?is)<style[\s\S]*?</style>", " ", html)
    html = re.sub(r"(?is)<[^>]+>", "\n", html)
    text = unescape(html)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def title_tokens(title: str) -> list[str]:
    t = re.sub(r"[《》〈〉（）()〔〕\[\]\s、，。：:；;·\-—]", "", str(title or ""))
    toks = re.findall(r"[\u4e00-\u9fff]{2,8}", t)
    return toks[:6]


def text_matches(title: str, body: str) -> bool:
    if len(body) < 180:
        return False
    toks = title_tokens(title)
    if not toks:
        return False
    hit = sum(1 for tok in toks if tok in body)
    return hit >= min(2, len(toks))


def fetch_shidian_chapter(url: str) -> str:
    m = re.search(r"/book/([^/]+)/chapter/([^/?#]+)", url)
    if not m:
        return ""
    book_id, cid = m.group(1), m.group(2)
    html = fetch_html(url, CACHE / "shidian" / f"{book_id}_{cid}.html")
    payload = book_payload(load_router(html))
    paragraphs = payload.get("paragraphList") or []
    return "\n".join(t for t in (paragraph_text(p).strip() for p in paragraphs) if t)


def fetch_wikisource(url: str) -> str:
    path = unquote(urlparse(url).path)
    if "/wiki/" not in path:
        return ""
    title = path.split("/wiki/", 1)[1]
    if not title or title.endswith("/"):
        return ""
    return get_wikitext(title) or ""


def fetch_generic(url: str) -> str:
    r = SESSION.get(url, timeout=40)
    r.raise_for_status()
    r.encoding = r.apparent_encoding or "utf-8"
    return html_to_text(r.text)


def extract_from_url(url: str) -> str:
    if not url or not url.startswith("http"):
        return ""
    if "ctext.org" in url:
        return ""
    try:
        if "shidianguji.com" in url and "/chapter/" in url:
            return fetch_shidian_chapter(url)
        if "wikisource.org" in url:
            return fetch_wikisource(url)
        if host_ok(url):
            return fetch_generic(url)
    except Exception as exc:
        log_event({"event": "fetch_fail", "url": url, "error": str(exc)[:200]})
    return ""


def dump_found(seq, title: str, url: str, body: str, extra: dict | None = None) -> None:
    extra = extra or {}
    header = f"# {title}\n# 来源: {url}\n# 编码: UTF-8\n# 归档日期: {STAMP}\n\n"
    content = header + body.strip() + "\n"
    folder = RESULT / folder_name(seq, title)
    product = PRODUCT / str(title).strip("《》")
    folder.mkdir(parents=True, exist_ok=True)
    product.mkdir(parents=True, exist_ok=True)
    fname = "正文.txt"
    (folder / fname).write_text(content, encoding="utf-8")
    (product / fname).write_text(content, encoding="utf-8")
    (folder / "来源网址.txt").write_text(url + "\n", encoding="utf-8")
    (product / "来源网址.txt").write_text(url + "\n", encoding="utf-8")
    meta = {
        "seq": seq,
        "title": title,
        "url": url,
        "publisher": extra.get("publisher") or "不详",
        "edition": extra.get("edition") or "不详",
        "note": extra.get("note") or f"{STAMP} 已核验公开录文并归档。",
        "saved_at": STAMP,
        "ok": True,
        "chars": len(body),
    }
    (folder / "基本信息.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    (product / "基本信息.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")


def mark_excel_found(found: list[dict]) -> int:
    if not found:
        return 0
    wb = load_workbook(EXCEL)
    ws = wb[wb.sheetnames[0]]
    by = {int(it["seq"]): it for it in found}
    n = 0
    for row in ws.iter_rows(min_row=5):
        if row[0].value != OWNER:
            continue
        seq = row[1].value
        if seq not in by:
            continue
        it = by[seq]
        row[5].value = it["url"]
        row[6].value = "是"
        if it.get("publisher"):
            row[7].value = it["publisher"]
        if it.get("edition"):
            row[8].value = it["edition"]
        old = str(row[9].value or "").strip()
        note = it.get("note") or f"{STAMP} 已核验公开录文并归档。"
        row[9].value = note if old in ("", "未找到", "无") else f"{old}；{note}"
        n += 1
    save_wb(wb, EXCEL)
    return n


def try_verify(plans: list, extras: list[dict] | None = None) -> list[dict]:
    found: list[dict] = []
    seen: set[int] = set()
    candidates: list[tuple] = []
    for p in plans or []:
        seq = p.get("seq")
        title = p.get("title") or ""
        for u in p.get("candidate_urls") or []:
            url = (u or {}).get("url") or ""
            likely = bool((u or {}).get("likely_fulltext"))
            if url and host_ok(url):
                candidates.append((seq, title, url, likely))
    for rec in extras or []:
        url = str(rec.get("url") or "")
        if url.startswith("http") and host_ok(url):
            candidates.append((rec["seq"], rec["title"], url, True))

    for seq, title, url, likely in candidates:
        try:
            seq_i = int(seq)
        except (TypeError, ValueError):
            continue
        if seq_i in seen:
            continue
        body = extract_from_url(url)
        if not body:
            continue
        if not likely and len(body) < 400:
            continue
        if not text_matches(str(title), body):
            log_event(
                {
                    "event": "verify_mismatch",
                    "seq": seq_i,
                    "url": url,
                    "chars": len(body),
                }
            )
            continue
        note = f"{STAMP} 核验 Astra/原表候选页，已保存匹配录文。"
        dump_found(seq_i, str(title).strip("《》"), url, body, {"note": note})
        found.append(
            {
                "seq": seq_i,
                "title": str(title).strip("《》"),
                "url": url,
                "note": note,
                "publisher": "不详",
                "edition": "公开录文页",
            }
        )
        seen.add(seq_i)
        log_event({"event": "archived", "seq": seq_i, "title": title, "chars": len(body), "url": url})
        time.sleep(0.4)
    n = mark_excel_found(found)
    if n:
        log_event({"event": "excel_found", "n": n, "seqs": [f["seq"] for f in found]})
    return found


def count_status() -> tuple[int, int]:
    wb = load_workbook(EXCEL, data_only=True)
    ws = wb[wb.sheetnames[0]]
    yes = no = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] != OWNER or row[2] is None:
            continue
        if str(row[6]) == "是":
            yes += 1
        else:
            no += 1
    return yes, no


def update_summary(reason: str, extra_found: int) -> None:
    yes, no = count_status()
    line = (
        f"{STAMP} 额度循环：{reason}。本进程新归档 {extra_found} 条。"
        f"当前找到{yes}条，未找到{no}条。\n"
    )
    old = SUMMARY.read_text(encoding="utf-8") if SUMMARY.exists() else ""
    if "共154条" in old:
        old = re.sub(r"找到\d+条，未找到\d+条。", f"找到{yes}条，未找到{no}条。", old, count=1)
    SUMMARY.write_text(line + old, encoding="utf-8")
    log_event({"event": "summary", "yes": yes, "no": no, "reason": reason, "extra_found": extra_found})


def is_quota(exc: BaseException) -> bool:
    if isinstance(exc, LlmQuotaExhaustedError) or is_llm_quota_exhausted_error(exc):
        return True
    text = f"{type(exc).__name__}: {exc}".lower()
    markers = (
        "insufficient_user_quota",
        "额度不足",
        "剩余额度",
        "quota_exhausted",
        "exceeded your current quota",
        "quota exceeded",
        "余额不足",
        "credit",
        "billing",
    )
    return any(m.lower() in text for m in markers)


def main() -> int:
    total_found = 0
    batches = 0
    deep_rounds = 0
    log_event({"event": "start", "owner": OWNER, "batch_size": BATCH_SIZE})
    leftover = load_remaining(only_unplanned=False)
    pre = [r for r in leftover if str(r.get("url") or "").startswith("http") and host_ok(str(r["url"]))]
    if pre:
        got = try_verify([], extras=pre)
        total_found += len(got)

    try:
        while True:
            unplanned = prioritize(load_remaining(only_unplanned=True))
            if unplanned:
                head = USER_HEAD
                batch = unplanned[:BATCH_SIZE]
            else:
                deep_rounds += 1
                pool = prioritize(load_remaining(only_unplanned=False))
                if not pool:
                    update_summary("剩余条目已全部处理", total_found)
                    return 0
                start = ((deep_rounds - 1) * BATCH_SIZE) % max(len(pool), 1)
                batch = pool[start : start + BATCH_SIZE]
                if len(batch) < BATCH_SIZE:
                    batch += pool[: BATCH_SIZE - len(batch)]
                head = DEEP_HEAD
                if deep_rounds > 300:
                    update_summary("深挖轮次上限", total_found)
                    return 0
            batches += 1
            print(f"\n=== batch {batches} seqs={[b['seq'] for b in batch]} deep={deep_rounds} ===", flush=True)
            try:
                plans = call_astra(batch, head)
            except Exception as exc:
                if is_quota(exc):
                    log_event({"event": "quota_exhausted", "error": str(exc)[:400], "batches": batches})
                    update_summary("gpt-6-astra 额度耗尽", total_found)
                    return 73
                log_event({"event": "astra_error", "error": f"{type(exc).__name__}: {exc}"[:400]})
                time.sleep(8)
                continue
            try:
                apply_plans(plans)
            except Exception as exc:
                log_event({"event": "apply_error", "error": f"{type(exc).__name__}: {exc}"[:400]})
            extras = [b for b in batch if str(b.get("url") or "").startswith("http")]
            try:
                got = try_verify(plans, extras=extras)
            except Exception as exc:
                log_event({"event": "verify_error", "error": f"{type(exc).__name__}: {exc}"[:400]})
                got = []
            total_found += len(got)
            yes, no = count_status()
            log_event({"event": "progress", "yes": yes, "no": no, "batches": batches, "new": len(got)})
            time.sleep(SLEEP_BETWEEN)
    except KeyboardInterrupt:
        update_summary("人工中断", total_found)
        return 130


if __name__ == "__main__":
    raise SystemExit(main())
