# -*- coding: utf-8 -*-
"""归档康熙三十六年出幸榆林：母项《榆林府志》艺文志御制诗。"""
from __future__ import annotations

import json
import sys
import time
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(r"D:\AI_Project\ancient_files_crawler")
RESULT = ROOT / "韦广波" / "爬取结果"
PRODUCT = ROOT / "韦广波" / "成果"
EXCEL = ROOT / "第22组总览表.xlsx"
STAMP = datetime.now().strftime("%Y-%m-%d")

SEQ = 8557
TITLE = "康熙三十六年康熙帝出幸榆林时对榆林的认识"
FOLDER = "8557康熙三十六年康熙帝出幸榆林时对榆林的认识"
FILE = "榆林府志_艺文志_御制河套西望_出塞.txt"
URL = "https://www.shidianguji.com/book/NLG312001065669"
BODY = """# 母项：清道光《榆林府志》艺文志·御制诗
# 子项：康熙三十六年圣祖平噶尔丹出幸榆林时所录御制二首

御製河套西望：康熙三十六年，居庸關由大同。聖祖仁皇帝平噶爾丹，出幸榆林。謹錄御製。二首。
往代存虛議，今爲我外藩。河環沙磧暖，境濶草灘繁。錯落延綏接，迷離朔漠吞。時巡曾不到，特示撫柔恩。

御製出塞：沿邊山行，岡嶺崎嶇，黃沙彌望，因自榆林鎭取道塞外，沙淺途平，水草咸便，且十二曰之程減半，得達安邊堡。書以記之。
森森萬騎厯駝城，榆林鎭一沙，名駝城。塞風淸磧路平，冰冸長河堪飮馬。月來大野照移營。郵籖紀地旬餘驛，羽轡行邊六日程。天下一家無内外，烽銷堠罷不論兵。
"""
NOTE = (
    f"{STAMP} 条目为史事/认识子项。母项道光《榆林府志》艺文志御制诗已核："
    "康熙三十六年圣祖平噶尔丹出幸榆林，录《河套西望》《出塞》。已从本轮识典抽出的府志文本摘录，不替代专论。"
)


def dump() -> None:
    header = f"# {TITLE}\n# 来源: {URL}\n# 编码: UTF-8\n# 归档日期: {STAMP}\n\n"
    content = header + BODY.strip() + "\n"
    folder = RESULT / FOLDER
    product = PRODUCT / TITLE
    folder.mkdir(parents=True, exist_ok=True)
    product.mkdir(parents=True, exist_ok=True)
    (folder / FILE).write_text(content, encoding="utf-8")
    (product / FILE).write_text(content, encoding="utf-8")
    (folder / "来源网址.txt").write_text(URL + "\n", encoding="utf-8")
    (product / "来源网址.txt").write_text(URL + "\n", encoding="utf-8")
    meta = {
        "seq": SEQ,
        "title": TITLE,
        "url": URL,
        "publisher": "不详",
        "edition": "清道光《榆林府志》艺文志御制诗；康熙三十六年出幸榆林。",
        "note": NOTE,
        "saved_at": STAMP,
        "ok": True,
    }
    (folder / "基本信息.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    (product / "基本信息.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    print("saved files", SEQ)


def mark_excel() -> None:
    last = None
    for i in range(10):
        try:
            wb = load_workbook(EXCEL)
            ws = wb[wb.sheetnames[0]]
            for row in ws.iter_rows(min_row=5):
                if row[0].value != "韦广波":
                    continue
                if row[1].value != SEQ:
                    continue
                row[5].value = URL
                row[6].value = "是"
                row[7].value = "不详"
                row[8].value = "道光《榆林府志》艺文志御制诗；康熙三十六年出幸榆林。"
                old = str(row[9].value or "").strip()
                row[9].value = NOTE if old in ("", "未找到", "无") else f"{old}；{NOTE}"
            wb.save(EXCEL)
            print("excel ok")
            return
        except Exception as exc:
            last = exc
            print("excel retry", i, type(exc).__name__)
            time.sleep(1.5)
    raise last  # type: ignore[misc]


if __name__ == "__main__":
    dump()
    mark_excel()
