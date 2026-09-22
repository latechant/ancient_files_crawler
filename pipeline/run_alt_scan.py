# -*- coding: utf-8 -*-
"""后台迭代扫描：识典检索/金石母项/维基文库，核验全文后才标「是」。不走 Astra。"""
from __future__ import annotations

import json
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import quote, unquote, urlparse

import requests
from openpyxl import load_workbook

ROOT = Path(r"D:\AI_Project\ancient_files_crawler")
sys.path.insert(0, str(ROOT))

from pipeline.extract_shidian import (  # noqa: E402
    book_payload,
    fetch_html,
    flatten_chapters,
    load_router,
    paragraph_text,
)
from pipeline.extract_wikisource import get_wikitext  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")

EXCEL = ROOT / "第22组总览表.xlsx"
RESULT = ROOT / "韦广波" / "爬取结果"
PRODUCT = ROOT / "韦广波" / "成果"
LOG = ROOT / "pipeline" / "logs"
CACHE = ROOT / "_cache"
SUMMARY = ROOT / "韦广波" / "爬虫总结.txt"
RUN_LOG = LOG / "alt_scan.jsonl"
LOG.mkdir(parents=True, exist_ok=True)

STAMP = datetime.now().strftime("%Y-%m-%d")
SLEEP = 0.9
MAX_ROUNDS = 8
WIKI_TIMEOUT = 20

HIGH_RE = re.compile(
    r"墓志|墓誌|墓碑|造像|碑|敕|奏|摺|志铭|志銘|寄谕|寄諭|题为|题请|"
    r"府志|县志|縣誌|告示|宗谱|宗譜|公报|公報|金石"
)
STRICT_RE = re.compile(r"寄谕|寄諭|题为|题请|奏报|呈.*批令|相关公文")
SKIP_COPYRIGHT = re.compile(
    r"岳麓书院藏秦简|岭光电|忆往昔|五代十国墓志汇编|西夏李氏世谱|北京档案史料"
)
GENERIC_TITLES = {"执照", "方鼎", "人类", "10月27日", "行书屏"}

MOTHERS = [
    ("7549461577028599854", "1lcvz8ltnwv0q", "八琼室续编"),
    ("CADAL01026682", "1lmm82h8zayib", "山右石刻丛编"),
    ("7659974904619466767", "1m73u9ulbvynn", "金石三编"),
    ("NGJ892411999028946160395", "1lqd64lans6qi", "益都金石略"),
    ("7562223221123055679", "1mdaekhnan22b", "鲁迅辑校墓志"),
    ("7563224556832292891", "1lx2cd4f5r19v", "鲁迅辑校石刻"),
    ("7440825182010212371", "1klkhvhmlszle", "全唐文"),
]

SESSION = requests.Session()
SESSION.headers.update(
    {
        "User-Agent": "Mozilla/5.0 (compatible; ancient-files-research/1.0; +local archival)",
        "Accept-Language": "zh-CN,zh;q=0.9",
    }
)


def log_event(obj: dict) -> None:
    obj = {"ts": datetime.now().isoformat(timespec="seconds"), **obj}
    with RUN_LOG.open("a", encoding="utf-8") as f:
        f.write(json.dumps(obj, ensure_ascii=False) + "\n")
    print(json.dumps(obj, ensure_ascii=False), flush=True)


def folder_name(seq, title: str) -> str:
    t = str(title or "").strip().strip("《》")
    t = re.sub(r'[\\/:*?"<>|]', "_", t)
    return f"{seq}{t}"


def load_remaining() -> list[dict]:
    wb = load_workbook(EXCEL, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] != "韦广波" or row[2] is None:
            continue
        if str(row[6]) == "是":
            continue
        rows.append(
            {
                "seq": row[1],
                "title": row[2],
                "url": row[5],
                "publisher": row[7],
                "edition": row[8],
                "note": row[9],
            }
        )
    return rows


def prioritize(rows: list[dict]) -> list[dict]:
    def score(r: dict) -> tuple:
        title = str(r.get("title") or "")
        hi = 0 if HIGH_RE.search(title) else 1
        named = 0 if len(re.findall(r"[\u4e00-\u9fff]", title)) >= 6 else 1
        return (hi, named, int(r["seq"] or 0))

    return sorted(rows, key=score)


def title_tokens(title: str) -> list[str]:
    t = re.sub(r"[《》〈〉（）()〔〕\[\]\s、，。：:；;·\-—]", "", str(title or ""))
    toks = re.findall(r"[\u4e00-\u9fff]{2,8}", t)
    stop = {"相关公文", "情形折", "请旨事", "全文", "原文", "并序"}
    return [x for x in toks if x not in stop][:8]


def queries_for(title: str) -> list[str]:
    raw = str(title or "").strip().strip("《》")
    toks = title_tokens(raw)
    qs = [raw]
    if len(toks) >= 2:
        qs.append("".join(toks[:2]))
    for tok in toks:
        if len(tok) >= 3:
            qs.append(tok)
    out = []
    seen = set()
    for q in qs:
        if q and q not in seen:
            seen.add(q)
            out.append(q)
    return out[:4]


def text_matches(title: str, body: str) -> bool:
    t = str(title or "").strip().strip("《》")
    if t in GENERIC_TITLES:
        return False
    if SKIP_COPYRIGHT.search(t + body[:200]):
        return False
    if len(body) < 200:
        return False
    toks = title_tokens(title)
    if not toks or not any(len(tok) >= 3 for tok in toks):
        return False
    hit = sum(1 for tok in toks if tok in body)
    need = 3 if STRICT_RE.search(str(title) or "") else min(2, len(toks))
    return hit >= need and any(len(tok) >= 3 and tok in body for tok in toks)


def walk_hits(obj, acc: list[dict] | None = None) -> list[dict]:
    acc = acc if acc is not None else []
    if isinstance(obj, dict):
        book = obj.get("bookId") or obj.get("book_id")
        chap = obj.get("chapterId") or obj.get("chapter_id")
        name = obj.get("chapterName") or obj.get("bookName") or obj.get("title") or obj.get("name")
        url = obj.get("url") or obj.get("jumpUrl") or ""
        if book and chap:
            acc.append(
                {
                    "bookId": str(book),
                    "chapterId": str(chap),
                    "name": str(name or ""),
                    "url": f"https://www.shidianguji.com/book/{book}/chapter/{chap}",
                }
            )
        elif isinstance(url, str) and "/book/" in url and "/chapter/" in url:
            acc.append({"bookId": "", "chapterId": "", "name": str(name or ""), "url": url})
        for v in obj.values():
            walk_hits(v, acc)
    elif isinstance(obj, list):
        for it in obj:
            walk_hits(it, acc)
    return acc


def shidian_search(q: str) -> list[dict]:
    url = f"https://www.shidianguji.com/search?query={quote(q)}"
    safe = re.sub(r"[^\w\u4e00-\u9fff]+", "_", q)[:40]
    try:
        html = fetch_html(url, CACHE / "shidian_search" / f"{safe}.html")
        router = load_router(html)
    except Exception as exc:
        log_event({"event": "search_fail", "q": q[:40], "error": type(exc).__name__})
        return []
    hits = walk_hits(router)
    uniq = []
    seen = set()
    for h in hits:
        if h["url"] in seen:
            continue
        seen.add(h["url"])
        uniq.append(h)
    return uniq[:12]


def wiki_search(q: str) -> list[dict]:
    try:
        r = SESSION.get(
            "https://zh.wikisource.org/w/api.php",
            params={
                "action": "query",
                "list": "search",
                "srsearch": q,
                "srlimit": 8,
                "format": "json",
            },
            timeout=WIKI_TIMEOUT,
        )
        r.raise_for_status()
        return [
            {
                "title": it.get("title") or "",
                "url": "https://zh.wikisource.org/wiki/" + quote(it.get("title") or ""),
            }
            for it in r.json().get("query", {}).get("search", [])
        ]
    except Exception as exc:
        log_event({"event": "wiki_fail", "q": q, "error": type(exc).__name__})
        return []


def fetch_shidian_chapter(url: str) -> str:
    m = re.search(r"/book/([^/]+)/chapter/([^/?#]+)", url)
    if not m:
        return ""
    book_id, cid = m.group(1), m.group(2)
    html = fetch_html(url, CACHE / "shidian" / f"{book_id}_{cid}.html")
    payload = book_payload(load_router(html))
    paras = payload.get("paragraphList") or []
    return "\n".join(t for t in (paragraph_text(p).strip() for p in paras) if t)


def fetch_wiki_title(title: str) -> str:
    return get_wikitext(title) or ""


def save_wb(wb, path: Path, tries: int = 10) -> None:
    last = None
    for i in range(tries):
        try:
            wb.save(path)
            return
        except (OSError, PermissionError, EOFError) as exc:
            last = exc
            time.sleep(1.2 + i * 0.4)
    raise last  # type: ignore[misc]


def dump_found(seq, title: str, url: str, body: str, note: str) -> None:
    header = f"# {title}\n# 来源: {url}\n# 编码: UTF-8\n# 归档日期: {STAMP}\n\n"
    content = header + body.strip() + "\n"
    folder = RESULT / folder_name(seq, title)
    product = PRODUCT / str(title).strip("《》")
    folder.mkdir(parents=True, exist_ok=True)
    product.mkdir(parents=True, exist_ok=True)
    (folder / "正文.txt").write_text(content, encoding="utf-8")
    (product / "正文.txt").write_text(content, encoding="utf-8")
    (folder / "来源网址.txt").write_text(url + "\n", encoding="utf-8")
    (product / "来源网址.txt").write_text(url + "\n", encoding="utf-8")
    meta = {
        "seq": seq,
        "title": title,
        "url": url,
        "publisher": "不详",
        "edition": "公开录文页",
        "note": note,
        "saved_at": STAMP,
        "ok": True,
        "chars": len(body),
    }
    blob = json.dumps(meta, ensure_ascii=False, indent=2)
    (folder / "基本信息.json").write_text(blob, encoding="utf-8")
    (product / "基本信息.json").write_text(blob, encoding="utf-8")


def mark_excel(found: list[dict]) -> int:
    if not found:
        return 0
    wb = load_workbook(EXCEL)
    ws = wb[wb.sheetnames[0]]
    by = {int(it["seq"]): it for it in found}
    n = 0
    for row in ws.iter_rows(min_row=5):
        if row[0].value != "韦广波":
            continue
        seq = row[1].value
        if seq not in by:
            continue
        it = by[seq]
        row[5].value = it["url"]
        row[6].value = "是"
        row[7].value = it.get("publisher") or "不详"
        row[8].value = it.get("edition") or "公开录文页"
        old = str(row[9].value or "").strip()
        note = it.get("note") or f"{STAMP} 后台扫描核验归档。"
        row[9].value = note if old in ("", "未找到", "无") else f"{old}；{note}"
        n += 1
    save_wb(wb, EXCEL)
    return n


def count_status() -> tuple[int, int]:
    wb = load_workbook(EXCEL, data_only=True)
    ws = wb[wb.sheetnames[0]]
    yes = no = 0
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] != "韦广波" or row[2] is None:
            continue
        if str(row[6]) == "是":
            yes += 1
        else:
            no += 1
    return yes, no


def update_summary(extra: int, reason: str) -> None:
    yes, no = count_status()
    line = f"{STAMP} 后台异源扫描：{reason}。本进程新归档 {extra} 条。当前找到{yes}条，未找到{no}条。\n"
    old = SUMMARY.read_text(encoding="utf-8") if SUMMARY.exists() else ""
    if "共154条" in old:
        old = re.sub(r"找到\d+条，未找到\d+条。", f"找到{yes}条，未找到{no}条。", old, count=1)
    SUMMARY.write_text(line + old, encoding="utf-8")
    log_event({"event": "summary", "yes": yes, "no": no, "extra": extra, "reason": reason})


def try_archive(seq, title: str, url: str, body: str, seen: set[int]) -> dict | None:
    try:
        seq_i = int(seq)
    except (TypeError, ValueError):
        return None
    if seq_i in seen:
        return None
    if not text_matches(str(title), body):
        log_event({"event": "mismatch", "seq": seq_i, "url": url, "chars": len(body)})
        return None
    note = f"{STAMP} 后台异源扫描核验，已保存匹配录文。"
    dump_found(seq_i, str(title).strip("《》"), url, body, note)
    seen.add(seq_i)
    rec = {
        "seq": seq_i,
        "title": str(title).strip("《》"),
        "url": url,
        "note": note,
        "publisher": "不详",
        "edition": "公开录文页",
    }
    log_event({"event": "archived", "seq": seq_i, "title": title, "chars": len(body), "url": url})
    return rec


def scan_item(rec: dict, seen: set[int]) -> list[dict]:
    title = str(rec["title"])
    found = []
    if str(title).strip().strip("《》") in GENERIC_TITLES:
        log_event({"event": "skip_generic", "seq": rec["seq"]})
        return found
    if SKIP_COPYRIGHT.search(title + str(rec.get("note") or "")):
        log_event({"event": "skip_copyright", "seq": rec["seq"]})
        return found
    qs = queries_for(title)
    url0 = str(rec.get("url") or "")
    candidates: list[tuple[str, str]] = []
    if "shidianguji.com" in url0 and "/chapter/" in url0:
        candidates.append((url0, "shidian"))
    for q in qs:
        for h in shidian_search(q):
            candidates.append((h["url"], "shidian"))
        time.sleep(SLEEP)
        for h in wiki_search(q):
            candidates.append((h["title"], "wiki"))
        time.sleep(0.35)
    seen_u = set()
    for target, kind in candidates:
        if target in seen_u:
            continue
        seen_u.add(target)
        try:
            if kind == "shidian":
                body = fetch_shidian_chapter(target)
                url = target
            else:
                body = fetch_wiki_title(target)
                url = "https://zh.wikisource.org/wiki/" + quote(target)
        except Exception as exc:
            log_event({"event": "fetch_fail", "seq": rec["seq"], "error": type(exc).__name__, "target": target[:120]})
            continue
        if not body:
            continue
        hit = try_archive(rec["seq"], title, url, body, seen)
        if hit:
            found.append(hit)
            break
        time.sleep(0.25)
    return found


MOTHER_DONE = LOG / "alt_scan_mothers_done.txt"


def walk_mothers(rows: list[dict], seen: set[int]) -> list[dict]:
    if MOTHER_DONE.exists():
        log_event({"event": "mother_skip", "reason": "already done"})
        return []
    needles = []
    for rec in rows:
        for tok in title_tokens(str(rec["title"])):
            if len(tok) >= 3:
                needles.append((tok, rec))
    found = []
    for book_id, seed, label in MOTHERS:
        url = f"https://www.shidianguji.com/book/{book_id}/chapter/{seed}"
        try:
            html = fetch_html(url, CACHE / "shidian" / f"{book_id}_{seed}.html")
            payload = book_payload(load_router(html))
            info = payload["bookInfo"]
            chapters = flatten_chapters(info.get("catalog", {}).get("chapters") or [])
        except Exception as exc:
            log_event({"event": "mother_fail", "book": label, "error": type(exc).__name__})
            continue
        log_event({"event": "mother", "book": label, "chapters": len(chapters)})
        for ch in chapters:
            name = ch.get("chapterName") or ""
            cid = ch.get("chapterId")
            if not cid or not ch.get("hasMainContent"):
                continue
            matched = [rec for tok, rec in needles if tok in name]
            if not matched and label in ("全唐文", "册府元龟"):
                continue
            if not matched:
                # 卷级目录不含子目时，抽正文再扫
                matched = rows
            chap_url = f"https://www.shidianguji.com/book/{book_id}/chapter/{cid}"
            try:
                body = fetch_shidian_chapter(chap_url)
            except Exception:
                continue
            if not body:
                continue
            for rec in matched:
                if int(rec["seq"]) in seen:
                    continue
                if any(tok in body for tok in title_tokens(str(rec["title"]))[:3]):
                    hit = try_archive(rec["seq"], rec["title"], chap_url, body, seen)
                    if hit:
                        found.append(hit)
            time.sleep(SLEEP)
    MOTHER_DONE.write_text("done\n", encoding="utf-8")
    return found


def main() -> None:
    log_event({"event": "start", "max_rounds": MAX_ROUNDS})
    total_new = 0
    seen: set[int] = set()
    for rnd in range(1, MAX_ROUNDS + 1):
        rows = prioritize(load_remaining())
        yes, no = count_status()
        log_event({"event": "round", "n": rnd, "remaining": len(rows), "yes": yes, "no": no})
        if not rows:
            break
        batch_found = []
        if rnd == 1:
            batch_found.extend(walk_mothers(rows, seen))
        for rec in rows:
            if int(rec["seq"]) in seen:
                continue
            try:
                batch_found.extend(scan_item(rec, seen))
            except Exception as exc:
                log_event({"event": "item_fail", "seq": rec["seq"], "error": type(exc).__name__})
            time.sleep(0.2)
        n = mark_excel(batch_found)
        total_new += len(batch_found)
        log_event(
            {
                "event": "round_done",
                "n": rnd,
                "new": len(batch_found),
                "excel": n,
                "seqs": [x["seq"] for x in batch_found],
            }
        )
        if batch_found:
            update_summary(len(batch_found), f"第{rnd}轮")
        time.sleep(8)
    update_summary(total_new, "扫描结束")
    log_event({"event": "end", "new": total_new})


if __name__ == "__main__":
    main()
