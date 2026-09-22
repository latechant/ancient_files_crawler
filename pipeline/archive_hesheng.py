# -*- coding: utf-8 -*-
"""归档贺胜墓：母项虞集《道园学古录》卷十八《贺丞相墓志铭》。"""
from __future__ import annotations

import json
import sys
import time
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(r"D:\AI_Project\ancient_files_crawler")
RESULT = ROOT / "韦广波" / "爬取结果"
PRODUCT = ROOT / "韦广波" / "成果"
EXCEL = ROOT / "第22组总览表.xlsx"
STAMP = datetime.now().strftime("%Y-%m-%d")

SEQ = 8941
TITLE = "贺胜墓"
FOLDER = "8941贺胜墓"
FILE = "道园学古录_卷十八_贺丞相墓志铭.txt"
URL = "https://www.shidianguji.com/book/SK2783/chapter/1kf4dcej15q1m"
BODY = (ROOT / "pipeline" / "logs" / "_hesheng.txt").read_text(encoding="utf-8")
NOTE = (
    f"{STAMP} 条目为贺胜墓。母项虞集《道园学古录》卷十八《贺丞相墓志铭》已从识典四库本抽出全文。"
    "墓主贺胜字贞卿，小字伯颜/巴延，泰定追封秦国公谥惠愍。拓片页另见书法网，不作正文。"
)


def dump() -> None:
    header = f"# {TITLE}\n# 来源: {URL}\n# 编码: UTF-8\n# 归档日期: {STAMP}\n\n"
    content = header + BODY.strip() + "\n"
    folder = RESULT / FOLDER
    product = PRODUCT / TITLE
    folder.mkdir(parents=True, exist_ok=True)
    product.mkdir(parents=True, exist_ok=True)
    (folder / FILE).write_text(content, encoding="utf-8")
    (product / FILE).write_text(content, encoding="utf-8")
    (folder / "来源网址.txt").write_text(URL + "\n", encoding="utf-8")
    (product / "来源网址.txt").write_text(URL + "\n", encoding="utf-8")
    meta = {
        "seq": SEQ,
        "title": TITLE,
        "url": URL,
        "publisher": "不详",
        "edition": "元虞集《道园学古录》卷十八；四库本识典古籍。",
        "note": NOTE,
        "saved_at": STAMP,
        "ok": True,
    }
    (folder / "基本信息.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    (product / "基本信息.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    print("saved files", SEQ, len(BODY))


def mark_excel() -> None:
    last = None
    for i in range(10):
        try:
            wb = load_workbook(EXCEL)
            ws = wb[wb.sheetnames[0]]
            for row in ws.iter_rows(min_row=5):
                if row[0].value != "韦广波":
                    continue
                if row[1].value != SEQ:
                    continue
                row[5].value = URL
                row[6].value = "是"
                row[7].value = "不详"
                row[8].value = "元虞集《道园学古录》卷十八《贺丞相墓志铭》；识典四库本。"
                old = str(row[9].value or "").strip()
                row[9].value = NOTE if old in ("", "未找到", "无") else f"{old}；{NOTE}"
            wb.save(EXCEL)
            print("excel ok")
            return
        except Exception as exc:
            last = exc
            print("excel retry", i, type(exc).__name__)
            time.sleep(1.5)
    raise last  # type: ignore[misc]


if __name__ == "__main__":
    dump()
    mark_excel()
