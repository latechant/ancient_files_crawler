# -*- coding: utf-8 -*-
"""把 gpt-6-astra 检索方案写入总览表备注，不把首页线索改成「是」。"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(r"D:\AI_Project\ancient_files_crawler")
EXCEL = ROOT / "第22组总览表.xlsx"
PLANS = json.loads((ROOT / "pipeline" / "logs" / "astra_search_latest.json").read_text(encoding="utf-8"))
STAMP = datetime.now().strftime("%Y-%m-%d")

# 先前人工核到的比 Astra 首页更具体的出处，仍不算全文。
EXTRA = {
    2144: "人工核验：安徽省文物考古研究所、潜山县文物管理所《安徽潜山彭岭战国西汉墓》，《考古学报》2006年第2期。现代期刊，未取得可公开全文。",
    196: "人工核验：目标本为民国三十二年铅印本（曾枢修、凌开蔚纂，二十卷）；康熙/乾隆本另见书目，公开全文仍未取得。",
}

wb = load_workbook(EXCEL)
ws = wb[wb.sheetnames[0]]
by_seq = {int(p["seq"]): p for p in PLANS}
n = 0
for row in ws.iter_rows(min_row=5):
    if row[0].value != "韦广波":
        continue
    seq = row[1].value
    if seq not in by_seq:
        continue
    p = by_seq[seq]
    kws = "、".join((p.get("keywords") or [])[:6])
    urls = "; ".join(u.get("url") for u in (p.get("candidate_urls") or [])[:3])
    note = (
        f"{STAMP} gpt-6-astra：{p.get('verdict')}。"
        f"类型={p.get('item_type')}。关键词={kws}。"
        f"母项={('、'.join(p.get('mother_works') or [])[:80])}。"
        f"候选站={urls}。{p.get('notes') or ''}"
    )
    if seq in EXTRA:
        note += " " + EXTRA[seq]
    old = str(row[9].value or "").strip()
    row[9].value = note if old in ("", "未找到", "无") else f"{old}；{note}"
    # keep 否
    n += 1
    folder = ROOT / "韦广波" / "爬取结果" / f"{seq}{str(p['title']).strip('《》')}"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "本轮补充").mkdir(exist_ok=True)
    (folder / "本轮补充" / "astra检索方案.json").write_text(
        json.dumps(p, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (folder / "来源网址.txt").write_text(
        "\n".join(u.get("url") for u in (p.get("candidate_urls") or []) if u.get("url")) + "\n",
        encoding="utf-8",
    )

wb.save(EXCEL)
print("annotated", n)
